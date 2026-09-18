from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNS = [
    "OBRA",
    "TIPO_FOTO",
    "LINK_IDENTIFICADO_OCR",
    "LATITUDE",
    "LONGITUDE",
    "ARQUIVO_FOTO",
    "STATUS_LINK",
    "STATUS_COORDENADA",
    "OCR_RODAPE",
    "DATA_HORA",
]


def _style_excel(path: Path, sheet_name: str = "Coordenadas") -> None:
    wb = load_workbook(path)
    ws = wb[sheet_name]
    header_fill = PatternFill("solid", fgColor="0D256C")
    header_font = Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName="TabelaCoordenadas", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        tab.tableStyleInfo = style
        try:
            ws.add_table(tab)
        except ValueError:
            pass

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
            for c in cell:
                max_len = max(max_len, len(str(c.value or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(12, max_len + 2), 65)

    wb.save(path)


def save_work_excel(records: Iterable[dict], path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(list(records))
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[COLUMNS]
    df.to_excel(path, index=False, sheet_name="Coordenadas")
    _style_excel(path)
    return path


def update_consolidated(records: Iterable[dict], path: Path | str) -> Path:
    path = Path(path)
    new_df = pd.DataFrame(list(records))
    for col in COLUMNS:
        if col not in new_df.columns:
            new_df[col] = ""
    new_df = new_df[COLUMNS]

    if path.exists():
        try:
            old = pd.read_excel(path, sheet_name="Coordenadas")
        except Exception:
            old = pd.DataFrame(columns=COLUMNS)
        for col in COLUMNS:
            if col not in old.columns:
                old[col] = ""
        # Remove a versão anterior das mesmas obra/tipo para manter consolidado limpo.
        keys = set(zip(new_df["OBRA"].astype(str), new_df["TIPO_FOTO"].astype(str)))
        mask_keep = [
            (str(o), str(t)) not in keys
            for o, t in zip(old["OBRA"].astype(str), old["TIPO_FOTO"].astype(str))
        ]
        old = old.loc[mask_keep, COLUMNS]
        df = pd.concat([old, new_df], ignore_index=True)
    else:
        df = new_df

    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(path, index=False, sheet_name="Coordenadas")
    _style_excel(path)
    return path
