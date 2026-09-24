from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import load_workbook
from PIL import Image as PILImage
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import APP_DIR, BUNDLE_DIR

BASE_FILENAME = "BASE_LEVANTAMENTO_ATUALIZADA.xlsx"
BASE_SHEET = "NOTAS"


@dataclass
class WorkMetadata:
    protocolo: str = ""
    nome: str = ""
    conta_contrato: str = ""
    instalacao: str = ""
    fase: str = ""
    endereco: str = ""
    localidade: str = ""
    municipio: str = ""
    tipo_nota: str = ""
    informacoes_extras: str = ""


def _clean_number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    if text.endswith(".0"):
        maybe = text[:-2]
        if maybe.replace("-", "").isdigit():
            return maybe
    return text


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split()).strip()


def _truncate(text: str, limit: int = 520) -> str:
    text = _clean_text(text)
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 3)].rstrip() + "..."


def find_base_workbook() -> Optional[Path]:
    """Localiza a base de levantamento.

    Prioridade:
    1. copia atualizada em APP_DIR (gravavel no PC do usuario);
    2. arquivo empacotado no EXE;
    3. qualquer BASE_LEVANTAMENTO*.xlsx nas duas pastas.
    """
    direct = [
        APP_DIR / BASE_FILENAME,
        BUNDLE_DIR / BASE_FILENAME,
    ]
    for path in direct:
        if path.exists():
            return path

    for folder in (APP_DIR, BUNDLE_DIR):
        try:
            matches = sorted(folder.glob("BASE_LEVANTAMENTO*.xlsx"))
        except Exception:
            matches = []
        if matches:
            return matches[0]
    return None


@lru_cache(maxsize=4)
def _load_metadata_index_cached(path_str: str, mtime_ns: int) -> dict[str, WorkMetadata]:
    path = Path(path_str)
    wb = load_workbook(path, read_only=True, data_only=True)
    if BASE_SHEET not in wb.sheetnames:
        raise RuntimeError(f"A planilha '{BASE_SHEET}' nao foi encontrada em {path.name}.")

    ws = wb[BASE_SHEET]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {str(v).strip().upper(): idx for idx, v in enumerate(header_row) if v is not None}

    required = [
        "PROTOCOLO",
        "NOME",
        "CONTA CONTRATO",
        "INSTALAÇÃO",
        "FASE",
        "ENDEREÇO",
        "LOCALIDADE",
        "MUNICIPIO",
        "TIPO NOTA",
        "INFORMAÇÕES EXTRAS",
    ]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError("Colunas ausentes na base: " + ", ".join(missing))

    index: dict[str, WorkMetadata] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        protocolo = _clean_number(row[headers["PROTOCOLO"]])
        if not protocolo:
            continue
        index[protocolo] = WorkMetadata(
            protocolo=protocolo,
            nome=_clean_text(row[headers["NOME"]]),
            conta_contrato=_clean_number(row[headers["CONTA CONTRATO"]]),
            instalacao=_clean_number(row[headers["INSTALAÇÃO"]]),
            fase=_clean_number(row[headers["FASE"]]),
            endereco=_clean_text(row[headers["ENDEREÇO"]]),
            localidade=_clean_text(row[headers["LOCALIDADE"]]),
            municipio=_clean_text(row[headers["MUNICIPIO"]]),
            tipo_nota=_clean_text(row[headers["TIPO NOTA"]]),
            informacoes_extras=_truncate(row[headers["INFORMAÇÕES EXTRAS"]], 320),
        )
    wb.close()
    return index


def load_metadata_index(path: Optional[Path] = None) -> tuple[dict[str, WorkMetadata], Optional[Path]]:
    base = Path(path) if path else find_base_workbook()
    if not base or not base.exists():
        return {}, None
    stat = base.stat()
    return _load_metadata_index_cached(str(base.resolve()), int(stat.st_mtime_ns)), base


def _fmt_coord(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        num = float(value)
    except Exception:
        return str(value)
    return f"{num:.7f}".rstrip("0").rstrip(".")


def _add_hyperlink(paragraph, url: str, text: Optional[str] = None):
    text = text or url
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _set_default_font(doc: Document, name: str = "Arial", size: float = 10.5) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = name
    normal.font.size = Pt(size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _configure_section(section) -> None:
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)


def _add_field_line(doc: Document, label: str, value: str = "", font_size: float = 10.5) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run(label)
    r1.bold = False
    r1.font.name = "Arial"
    r1.font.size = Pt(font_size)
    r2 = p.add_run(value or "")
    r2.font.name = "Arial"
    r2.font.size = Pt(font_size)


def _fit_image_size(path: Path, max_w_cm: float = 9.2, max_h_cm: float = 10.0) -> tuple[Cm, Cm]:
    with PILImage.open(path) as im:
        w, h = im.size
    if w <= 0 or h <= 0:
        return Cm(max_w_cm), Cm(max_h_cm)
    ratio = w / h
    max_w = max_w_cm
    max_h = max_h_cm
    width = max_w
    height = width / ratio
    if height > max_h:
        height = max_h
        width = height * ratio
    return Cm(width), Cm(height)


def _resolve_facade_image(result: dict) -> Optional[Path]:
    obra = str(result.get("obra", ""))
    folder = Path(result.get("folder") or "")
    preferred = folder / f"{obra}_FACHADADOIMOVEL.jpg"
    if preferred.exists():
        return preferred

    for rec in result.get("records", []) or []:
        tipo = str(rec.get("TIPO_FOTO", "")).upper()
        if "FACHADA" not in tipo:
            continue
        raw = str(rec.get("ARQUIVO_FOTO", "")).strip()
        if raw:
            p = Path(raw)
            if p.exists():
                return p
    return None


def _metadata_for(index: dict[str, WorkMetadata], obra: str) -> WorkMetadata:
    return index.get(str(obra), WorkMetadata(protocolo=str(obra)))


def build_word_report(
    results: Iterable[dict],
    output_path: Path | str,
    base_path: Optional[Path | str] = None,
) -> tuple[Path, Optional[Path]]:
    """Gera um unico Word com uma obra por pagina.

    A coordenada usada no link Google Maps vem do OCR da FACHADA DO IMOVEL.
    Os demais dados sao obtidos na aba NOTAS da base, procurando PROTOCOLO.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    metadata_index, located_base = load_metadata_index(Path(base_path) if base_path else None)
    rows = [dict(r) for r in results if r.get("obra")]

    doc = Document()
    _set_default_font(doc)
    _configure_section(doc.sections[0])

    for idx, result in enumerate(rows):
        if idx > 0:
            doc.add_page_break()

        obra = str(result.get("obra", ""))
        meta = _metadata_for(metadata_index, obra)
        lat = result.get("fachada_latitude")
        lon = result.get("fachada_longitude")

        address = meta.endereco
        if meta.localidade:
            address = f"{address} / {meta.localidade}" if address else meta.localidade

        _add_field_line(doc, "NOTA CCS: ", meta.protocolo or obra)
        _add_field_line(doc, "NOME DO CLIENTE: ", meta.nome)
        _add_field_line(doc, "CONTA CONTRATO: ", meta.conta_contrato)
        _add_field_line(doc, "INSTALAÇÃO: ", meta.instalacao)
        _add_field_line(doc, "FASE: ", meta.fase)
        _add_field_line(doc, "ENDEREÇO: ", address)
        _add_field_line(doc, "MUNICÍPIO: ", meta.municipio)
        _add_field_line(doc, "INFORMAÇÕES EXTRAS: ", meta.informacoes_extras, font_size=9.0 if meta.informacoes_extras else 10.5)
        _add_field_line(doc, "TIPO NOTA: ", meta.tipo_nota)

        p_map = doc.add_paragraph()
        p_map.paragraph_format.space_before = Pt(2)
        p_map.paragraph_format.space_after = Pt(3)
        p_map.paragraph_format.line_spacing = 1.0
        lat_s = _fmt_coord(lat)
        lon_s = _fmt_coord(lon)
        if lat_s and lon_s:
            url = f"https://www.google.com.br/maps/place/{lat_s},{lon_s}"
            _add_hyperlink(p_map, url, url)
        else:
            run = p_map.add_run("COORDENADA DA FACHADA NÃO RECONHECIDA")
            run.font.name = "Arial"
            run.font.size = Pt(10.0)
            run.font.color.rgb = RGBColor(192, 0, 0)

        photo = _resolve_facade_image(result)
        if photo and photo.exists():
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(0)
            p_img.alignment = WD_ALIGN_PARAGRAPH.LEFT
            width, height = _fit_image_size(photo)
            run = p_img.add_run()
            run.add_picture(str(photo), width=width, height=height)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run("FOTO FACHADA DO IMÓVEL NÃO LOCALIZADA")
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(192, 0, 0)

    if not rows:
        p = doc.add_paragraph("Nenhuma obra processada para gerar o relatório.")
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(11)

    doc.save(output_path)
    return output_path, located_base


def default_report_path(output_root: Path | str) -> Path:
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return root / f"RELATORIO_LEVANTAMENTO_{stamp}.docx"
